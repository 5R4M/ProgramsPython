import customtkinter as ctk
from docx import Document
from tkinter import filedialog, messagebox
import os

class ActaNotarialEditor(ctk.CTk):
    def __init__(self):
        super().__init__()

        self.title("Editor de Acta Notarial - 2025")
        ctk.set_appearance_mode("dark")
        ctk.set_default_color_theme("blue")

        self.plantilla_path = None
        self.plantilla_cargada = False

        # Definir tamaño de la ventana
        ancho_ventana = 1000
        alto_ventana = 900
        
        # Actualizar la ventana para obtener dimensiones reales de la pantalla
        self.update_idletasks()
        
        # Obtener dimensiones de la pantalla
        ancho_pantalla = self.winfo_screenwidth()
        alto_pantalla = self.winfo_screenheight()
        
        # Calcular posición centrada
        x = (ancho_pantalla - ancho_ventana) // 2
        y = (alto_pantalla - alto_ventana) // 2
        
        # Asegurar que y no sea negativo
        if y < 0:
            y = 0
        
        # Establecer geometría completa
        self.geometry(f"{ancho_ventana}x{alto_ventana}+{x}+{y}")
        
        # Forzar la actualización de la posición
        self.update()
        
        self.crear_interfaz()

    def crear_interfaz(self):
        main_frame = ctk.CTkScrollableFrame(self, width=850, height=750)
        main_frame.pack(pady=20, padx=20, fill="both", expand=True)

        titulo = ctk.CTkLabel(
            main_frame,
            text="📄 Editor de Acta Notarial",
            font=ctk.CTkFont(size=28, weight="bold")
        )
        titulo.pack(pady=20)

        self.btn_cargar = ctk.CTkButton(
            main_frame,
            text="📁 Cargar Plantilla (.docx)",
            command=self.cargar_plantilla,
            height=40,
            font=ctk.CTkFont(size=14)
        )
        self.btn_cargar.pack(pady=10)

        self.lbl_archivo = ctk.CTkLabel(
            main_frame,
            text="No se ha cargado ninguna plantilla",
            text_color="gray"
        )
        self.lbl_archivo.pack(pady=5)

        ctk.CTkLabel(main_frame, text="").pack(pady=5)

        # ----- Fecha y Hora -----
        frame_fecha = ctk.CTkFrame(main_frame)
        frame_fecha.pack(pady=10, padx=20, fill="x")

        ctk.CTkLabel(
            frame_fecha,
            text="🕐 Fecha y Hora",
            font=ctk.CTkFont(size=18, weight="bold")
        ).pack(pady=10)

        frame_hora = ctk.CTkFrame(frame_fecha)
        frame_hora.pack(pady=5, fill="x", padx=20)

        ctk.CTkLabel(frame_hora, text="Hora:", width=100).pack(side="left", padx=5)
        self.entry_hora = ctk.CTkEntry(frame_hora, placeholder_text="Ej: 17")
        self.entry_hora.pack(side="left", padx=5, expand=True, fill="x")

        ctk.CTkLabel(frame_hora, text="Minutos:", width=100).pack(side="left", padx=5)
        self.entry_minutos = ctk.CTkEntry(frame_hora, placeholder_text="Ej: 20")
        self.entry_minutos.pack(side="left", padx=5, expand=True, fill="x")

        frame_fecha_dia = ctk.CTkFrame(frame_fecha)
        frame_fecha_dia.pack(pady=5, fill="x", padx=20)

        ctk.CTkLabel(frame_fecha_dia, text="Día:", width=100).pack(side="left", padx=5)
        self.entry_dia = ctk.CTkEntry(frame_fecha_dia, placeholder_text="Ej: 28")
        self.entry_dia.pack(side="left", padx=5, expand=True, fill="x")

        ctk.CTkLabel(frame_fecha_dia, text="Mes:", width=100).pack(side="left", padx=5)
        self.combo_mes = ctk.CTkComboBox(
            frame_fecha_dia,
            values=["enero", "febrero", "marzo", "abril", "mayo", "junio",
                    "julio", "agosto", "septiembre", "octubre", "noviembre", "diciembre"]
        )
        self.combo_mes.set("noviembre")
        self.combo_mes.pack(side="left", padx=5, expand=True, fill="x")

        ctk.CTkLabel(frame_fecha_dia, text="Año:", width=100).pack(side="left", padx=5)
        self.entry_anio = ctk.CTkEntry(frame_fecha_dia, placeholder_text="Ej: 2025")
        self.entry_anio.pack(side="left", padx=5, expand=True, fill="x")

        # ----- Datos personales -----
        frame_datos = ctk.CTkFrame(main_frame)
        frame_datos.pack(pady=10, padx=20, fill="x")

        ctk.CTkLabel(
            frame_datos,
            text="👤 Datos Personales",
            font=ctk.CTkFont(size=18, weight="bold")
        ).pack(pady=10)

        frame_nombre = ctk.CTkFrame(frame_datos)
        frame_nombre.pack(pady=5, fill="x", padx=20)
        ctk.CTkLabel(frame_nombre, text="Nombre Completo:", width=150).pack(side="left", padx=5)
        self.entry_nombre = ctk.CTkEntry(
            frame_nombre,
            placeholder_text="Ej: Abner Aníbal Ajpop González"
        )
        self.entry_nombre.pack(side="left", padx=5, expand=True, fill="x")

        frame_edad = ctk.CTkFrame(frame_datos)
        frame_edad.pack(pady=5, fill="x", padx=20)
        ctk.CTkLabel(frame_edad, text="Edad:", width=150).pack(side="left", padx=5)
        self.entry_edad = ctk.CTkEntry(frame_edad, placeholder_text="Ej: 21")
        self.entry_edad.pack(side="left", padx=5, expand=True, fill="x")

        frame_estado = ctk.CTkFrame(frame_datos)
        frame_estado.pack(pady=5, fill="x", padx=20)
        ctk.CTkLabel(frame_estado, text="Estado Civil:", width=150).pack(side="left", padx=5)
        self.combo_estado = ctk.CTkComboBox(
            frame_estado,
            values=["soltero", "soltera", "casado", "casada",
                    "divorciado", "divorciada", "viudo", "viuda"]
        )
        self.combo_estado.set("soltero")
        self.combo_estado.pack(side="left", padx=5, expand=True, fill="x")

        frame_casada = ctk.CTkFrame(frame_datos)
        frame_casada.pack(pady=5, fill="x", padx=20)
        ctk.CTkLabel(frame_casada, text="Apellido de casada:", width=150).pack(side="left", padx=5)
        self.entry_casada = ctk.CTkEntry(
            frame_casada,
            placeholder_text="(Opcional) Ej: de López"
        )
        self.entry_casada.pack(side="left", padx=5, expand=True, fill="x")

        frame_nacionalidad = ctk.CTkFrame(frame_datos)
        frame_nacionalidad.pack(pady=5, fill="x", padx=20)
        ctk.CTkLabel(frame_nacionalidad, text="Nacionalidad:", width=150).pack(side="left", padx=5)
        self.entry_nacionalidad = ctk.CTkEntry(
            frame_nacionalidad,
            placeholder_text="Ej: guatemalteco / guatemalteca"
        )
        self.entry_nacionalidad.pack(side="left", padx=5, expand=True, fill="x")

        frame_nivel = ctk.CTkFrame(frame_datos)
        frame_nivel.pack(pady=5, fill="x", padx=20)
        ctk.CTkLabel(frame_nivel, text="Nivel Académico:", width=150).pack(side="left", padx=5)
        self.entry_nivel = ctk.CTkEntry(
            frame_nivel,
            placeholder_text="Ej: Bachiller en Ciencias y Letras con Orientación en Computación"
        )
        self.entry_nivel.pack(side="left", padx=5, expand=True, fill="x")

        frame_domicilio = ctk.CTkFrame(frame_datos)
        frame_domicilio.pack(pady=5, fill="x", padx=20)
        ctk.CTkLabel(frame_domicilio, text="Domicilio:", width=150).pack(side="left", padx=5)
        self.entry_domicilio = ctk.CTkEntry(
            frame_domicilio,
            placeholder_text="Ej: departamento de Guatemala"
        )
        self.entry_domicilio.pack(side="left", padx=5, expand=True, fill="x")

        # ----- DPI -----
        frame_dpi = ctk.CTkFrame(frame_datos)
        frame_dpi.pack(pady=5, fill="x", padx=20)
        ctk.CTkLabel(frame_dpi, text="DPI (CUI):", width=150).pack(side="left", padx=5)
        self.entry_dpi = ctk.CTkEntry(
            frame_dpi,
            placeholder_text="Ej: 2008 22829 0101"
        )
        self.entry_dpi.pack(side="left", padx=5, expand=True, fill="x")

        # ----- Botones -----
        frame_botones = ctk.CTkFrame(main_frame)
        frame_botones.pack(pady=20)

        self.btn_procesar = ctk.CTkButton(
            frame_botones,
            text="✅ Generar Documento",
            command=self.generar_documento,
            height=45,
            width=200,
            font=ctk.CTkFont(size=16, weight="bold"),
            fg_color="#2ecc71",
            hover_color="#27ae60"
        )
        self.btn_procesar.pack(side="left", padx=10)

        self.btn_limpiar = ctk.CTkButton(
            frame_botones,
            text="🔄 Limpiar Campos",
            command=self.limpiar_campos,
            height=45,
            width=200,
            font=ctk.CTkFont(size=16, weight="bold"),
            fg_color="#e74c3c",
            hover_color="#c0392b"
        )
        self.btn_limpiar.pack(side="left", padx=10)

    def cargar_plantilla(self):
        if self.plantilla_cargada:
            resp = messagebox.askyesno(
                "Plantilla ya cargada",
                "Ya hay una plantilla cargada.\n¿Desea cargar una nueva?"
            )
            if not resp:
                return

        archivo = filedialog.askopenfilename(
            title="Seleccionar plantilla Word",
            filetypes=[("Documento Word (.docx)", "*.docx")]
        )

        if not archivo:
            return

        if not archivo.lower().endswith(".docx"):
            messagebox.showerror(
                "Error",
                "El archivo seleccionado no es un .docx.\n"
                "Convierta primero su .doc a .docx en Word."
            )
            return

        try:
            _ = Document(archivo)
            self.plantilla_path = archivo
            self.plantilla_cargada = True
            nombre_archivo = os.path.basename(archivo)
            self.lbl_archivo.configure(
                text=f"✓ Plantilla cargada: {nombre_archivo}",
                text_color="green"
            )
            self.btn_cargar.configure(text="✓ Plantilla cargada", state="disabled")
            messagebox.showinfo(
                "Éxito",
                "Plantilla cargada correctamente.\n\n"
                "Ahora complete los campos y genere documentos."
            )
        except Exception as e:
            messagebox.showerror(
                "Error",
                f"No se pudo abrir el documento como .docx:\n{str(e)}"
            )

    def numero_a_texto(self, num: int) -> str:
        unidades = ["", "uno", "dos", "tres", "cuatro", "cinco",
                    "seis", "siete", "ocho", "nueve"]
        decenas = ["", "diez", "veinte", "treinta", "cuarenta",
                   "cincuenta", "sesenta", "setenta", "ochenta", "noventa"]
        especiales = ["diez", "once", "doce", "trece", "catorce",
                      "quince", "dieciséis", "diecisiete", "dieciocho", "diecinueve"]

        if num < 10:
            return unidades[num]
        if 10 <= num < 20:
            return especiales[num - 10]
        if 20 <= num < 30:
            if num == 20:
                return "veinte"
            return "veinti" + unidades[num - 20]
        if num < 100:
            d = num // 10
            u = num % 10
            if u == 0:
                return decenas[d]
            return f"{decenas[d]} y {unidades[u]}"
        return str(num)

    def convertir_dpi_a_texto(self, dpi: str) -> str:
        """
        Convierte un DPI en formato '2008 22829 0101' o '1916024260101' a texto.
        """
        dpi_limpio = dpi.replace(" ", "")
        
        if len(dpi_limpio) == 13 and dpi_limpio.isdigit():
            partes = [dpi_limpio[:4], dpi_limpio[4:9], dpi_limpio[9:13]]
        else:
            partes = dpi.strip().split()
            if len(partes) != 3:
                return dpi
        
        resultado = []
        
        for i, parte in enumerate(partes):
            if not parte.isdigit():
                resultado.append(parte)
                continue
            
            num = int(parte)
            
            if i == 0:
                if num >= 2000:
                    resto = num - 2000
                    if resto == 0:
                        resultado.append("dos mil")
                    else:
                        resultado.append("dos mil " + self.numero_a_texto(resto))
                elif num >= 1000:
                    miles = num // 1000
                    resto = num % 1000
                    texto_partes = []
                    if miles == 1:
                        texto_partes.append("mil")
                    else:
                        texto_partes.append(self.numero_a_texto(miles) + " mil")
                    if resto > 0:
                        texto_partes.append(self.numero_a_texto_centenas(resto))
                    resultado.append(" ".join(texto_partes))
                else:
                    resultado.append(self.numero_a_texto(num))
            
            elif i == 1:
                if parte.startswith("0") and len(parte) == 5:
                    resultado.append("cero " + self.convertir_numero_miles(int(parte[1:])))
                else:
                    resultado.append(self.convertir_numero_miles(num))
            
            elif i == 2:
                if parte.startswith("0") and len(parte) == 4:
                    resultado.append("cero " + self.numero_a_texto_centenas(int(parte[1:])))
                else:
                    if num == 0:
                        resultado.append("cero")
                    else:
                        resultado.append(self.convertir_numero_miles(num))
        
        return " espacio ".join(resultado)
    
    def convertir_numero_miles(self, num: int) -> str:
        if num < 1000:
            return self.numero_a_texto_centenas(num)
        
        miles = num // 1000
        resto = num % 1000
        
        texto_partes = []
        if miles == 1:
            texto_partes.append("mil")
        else:
            texto_partes.append(self.numero_a_texto(miles) + " mil")
        
        if resto > 0:
            texto_partes.append(self.numero_a_texto_centenas(resto))
        
        return " ".join(texto_partes)

    def numero_a_texto_centenas(self, num: int) -> str:
        if num < 100:
            return self.numero_a_texto(num)
        
        centenas_texto = ["", "ciento", "doscientos", "trescientos", "cuatrocientos",
                          "quinientos", "seiscientos", "setecientos", "ochocientos", "novecientos"]
        
        c = num // 100
        resto = num % 100
        
        if num == 100:
            return "cien"
        
        resultado = centenas_texto[c]
        if resto > 0:
            resultado += " " + self.numero_a_texto(resto)
        
        return resultado

    def reemplazar_en_parrafo(self, paragraph, buscar, reemplazar, reemplazar_todas=False):
        reemplazos_realizados = 0
        
        while True:
            texto_completo = ''.join(run.text for run in paragraph.runs)
            
            if buscar not in texto_completo:
                break
            
            pos_inicio = texto_completo.find(buscar)
            pos_fin = pos_inicio + len(buscar)
            
            pos_actual = 0
            run_inicio = None
            run_fin = None
            pos_en_run_inicio = 0
            pos_en_run_fin = 0
            
            for i, run in enumerate(paragraph.runs):
                len_run = len(run.text)
                
                if pos_actual <= pos_inicio < pos_actual + len_run:
                    run_inicio = i
                    pos_en_run_inicio = pos_inicio - pos_actual
                
                if pos_actual < pos_fin <= pos_actual + len_run:
                    run_fin = i
                    pos_en_run_fin = pos_fin - pos_actual
                    break
                
                pos_actual += len_run
            
            if run_inicio is None or run_fin is None:
                break
            
            if run_inicio == run_fin:
                run = paragraph.runs[run_inicio]
                run.text = run.text[:pos_en_run_inicio] + reemplazar + run.text[pos_en_run_fin:]
            else:
                paragraph.runs[run_inicio].text = paragraph.runs[run_inicio].text[:pos_en_run_inicio] + reemplazar
                
                runs_a_eliminar = []
                for i in range(run_inicio + 1, run_fin + 1):
                    if i == run_fin:
                        paragraph.runs[i].text = paragraph.runs[i].text[pos_en_run_fin:]
                        if not paragraph.runs[i].text:
                            runs_a_eliminar.append(i)
                    else:
                        runs_a_eliminar.append(i)
                
                for i in reversed(runs_a_eliminar):
                    paragraph._element.remove(paragraph.runs[i]._element)
            
            reemplazos_realizados += 1
            
            if not reemplazar_todas:
                break
        
        return reemplazos_realizados > 0

    def generar_documento(self):
        if not self.plantilla_cargada:
            messagebox.showwarning("Advertencia", "Primero debe cargar una plantilla")
            return

        if not self.entry_nombre.get().strip():
            messagebox.showwarning("Advertencia", "Debe ingresar al menos el nombre completo")
            return

        try:
            doc = Document(self.plantilla_path)

            hora = self.entry_hora.get().strip()
            minutos = self.entry_minutos.get().strip()
            dia = self.entry_dia.get().strip()
            mes = self.combo_mes.get().strip()
            anio = self.entry_anio.get().strip()
            nombre = self.entry_nombre.get().strip()
            edad = self.entry_edad.get().strip()
            estado_civil = self.combo_estado.get().strip()
            apellido_casada = self.entry_casada.get().strip()
            nacionalidad = self.entry_nacionalidad.get().strip()
            nivel_academico = self.entry_nivel.get().strip()
            domicilio = self.entry_domicilio.get().strip()
            dpi = self.entry_dpi.get().strip()

            nombre_completo = nombre
            if apellido_casada and estado_civil == "casada":
                partes = nombre.split()
                if len(partes) >= 2:
                    nombre_completo = f"{partes[0]} {partes[1]} {apellido_casada}"
                    if len(partes) > 2:
                        nombre_completo += " " + " ".join(partes[2:])

            anio_texto = None
            if anio.isdigit():
                anio_num = int(anio)
                if 2000 <= anio_num < 2100:
                    resto = anio_num - 2000
                    if resto == 0:
                        anio_texto = "dos mil"
                    else:
                        anio_texto = "dos mil " + self.numero_a_texto(resto)
                else:
                    anio_texto = str(anio)

            reemplazos = []

            if hora and minutos:
                hora_t = self.numero_a_texto(int(hora))
                min_t = self.numero_a_texto(int(minutos))
                reemplazos.append(("diecisiete horas con veinte minutos", f"{hora_t} horas con {min_t} minutos"))

            if dia:
                dia_t = self.numero_a_texto(int(dia))
                reemplazos.append(("veintiocho", dia_t))
                reemplazos.append(("(28)", f"({dia})"))

            if mes:
                reemplazos.append(("noviembre", mes))

            if anio and anio_texto:
                reemplazos.append(("dos mil veinticinco", anio_texto))
                reemplazos.append(("(2025)", f"({anio})"))

            if nombre:
                reemplazos.append(("Abner Aníbal Ajpop González", nombre_completo))

            if edad:
                edad_num = int(edad)
                edad_t = self.numero_a_texto(edad_num)
                
                if edad_t.endswith(" y uno"):
                    edad_t = edad_t[:-3] + " un"
                elif edad_t == "uno":
                    edad_t = "un"
                
                reemplazos.append(("veintiún", edad_t))
                reemplazos.append(("(21)", f"({edad})"))

            if estado_civil:
                reemplazos.append(("soltero", estado_civil))

            if nacionalidad:
                reemplazos.append(("guatemalteco", nacionalidad))

            if nivel_academico:
                reemplazos.append(("Bachiller en Ciencias y Letras con Orientación en Computación", nivel_academico))

            if domicilio:
                reemplazos.append(("con domicilio en el departamento de Guatemala", f"con domicilio en el {domicilio}"))

            if dpi:
                dpi_formateado = dpi.replace(" ", "")
                if len(dpi_formateado) == 13:
                    dpi_con_espacios = f"{dpi_formateado[:4]} {dpi_formateado[4:9]} {dpi_formateado[9:13]}"
                else:
                    dpi_con_espacios = dpi
                
                dpi_texto = self.convertir_dpi_a_texto(dpi)
                reemplazos.append(("dos mil ocho espacio veintidós mil ochocientos veintinueve espacio cero ciento uno", dpi_texto))
                reemplazos.append(("(2008 22829 0101)", f"({dpi_con_espacios})"))

            for paragraph in doc.paragraphs:
                for buscar, reemplazar in reemplazos:
                    if buscar == "Abner Aníbal Ajpop González":
                        self.reemplazar_en_parrafo(paragraph, buscar, reemplazar, reemplazar_todas=True)
                    else:
                        self.reemplazar_en_parrafo(paragraph, buscar, reemplazar, reemplazar_todas=False)

            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        for paragraph in cell.paragraphs:
                            for buscar, reemplazar in reemplazos:
                                if buscar == "Abner Aníbal Ajpop González":
                                    self.reemplazar_en_parrafo(paragraph, buscar, reemplazar, reemplazar_todas=True)
                                else:
                                    self.reemplazar_en_parrafo(paragraph, buscar, reemplazar, reemplazar_todas=False)

            archivo_salida = filedialog.asksaveasfilename(
                defaultextension=".docx",
                filetypes=[("Documento Word", "*.docx")],
                initialfile=f"acta_{nombre.replace(' ', '_')}.docx"
            )
            if archivo_salida:
                doc.save(archivo_salida)
                messagebox.showinfo(
                    "Éxito",
                    f"✅ Documento generado correctamente:\n\n{archivo_salida}"
                )

        except Exception as e:
            messagebox.showerror("Error", f"Error al generar el documento:\n{str(e)}")

    def limpiar_campos(self):
        self.entry_hora.delete(0, "end")
        self.entry_minutos.delete(0, "end")
        self.entry_dia.delete(0, "end")
        self.combo_mes.set("noviembre")
        self.entry_anio.delete(0, "end")
        self.entry_nombre.delete(0, "end")
        self.entry_edad.delete(0, "end")
        self.combo_estado.set("soltero")
        self.entry_casada.delete(0, "end")
        self.entry_nacionalidad.delete(0, "end")
        self.entry_nivel.delete(0, "end")
        self.entry_domicilio.delete(0, "end")
        self.entry_dpi.delete(0, "end")

if __name__ == "__main__":
    app = ActaNotarialEditor()
    app.mainloop()