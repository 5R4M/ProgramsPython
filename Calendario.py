from datetime import datetime, timedelta
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.cell.cell import Cell
from openpyxl.utils import get_column_letter

# Crear un nuevo archivo de Excel
wb = Workbook()
ws = wb.active

# Definir estilos y colores
header_fill = PatternFill(start_color="808080", end_color="808080", fill_type="solid")  # Gris oscuro
header_font = Font(color="FFFFFF", bold=True)  # Blanco para la fila de títulos
title_font = Font(color="000000", bold=True)  # Negro para los títulos del encabezado
title_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")  # Fondo blanco
special_fill = PatternFill(start_color="808080", end_color="808080", fill_type="solid")  # Gris oscuro
special_font = Font(color="FFFFFF", bold=True)  # Blanco
center_alignment = Alignment(horizontal="center", vertical="center")
border = Border(left=Side(style="thin"), right=Side(style="thin"),
                top=Side(style="thin"), bottom=Side(style="thin"))

# Definir los días festivos de Guatemala 2026
festivos_guatemala = {
    (1, 1): ('A', 'AÑO NUEVO'),
    (3, 29): ('A', 'DOMINGO DE RAMOS'),
    (4, 2): ('A', 'JUEVES SANTO'),
    (4, 3): ('A', 'VIERNES SANTO'),
    (4, 4): ('A', 'SÁBADO SANTO'),
    (4, 5): ('A', 'DOMINGO DE RESURRECCIÓN'),
    (5, 1): ('A', 'DÍA DEL TRABAJO'),
    (6, 30): ('A', 'DÍA DEL EJÉRCITO'),
    (8, 15): ('A', 'DÍA DE LA ASUNCIÓN'),
    (9, 15): ('A', 'DÍA DE LA INDEPENDENCIA'),
    (10, 20): ('A', 'DÍA DE LA REVOLUCIÓN'),
    (11, 1): ('A', 'DÍA DE TODOS LOS SANTOS'),
    (12, 24): ('A', 'NOCHEBUENA'),
    (12, 25): ('A', 'NAVIDAD'),
    (12, 31): ('A', 'FIN DE AÑO')
}

# Agregar los títulos del encabezado
titulos_encabezado = [
    "DIRECCIÓN DEPARTAMENTAL DE REDES INTEGRADAS DE SERVICIOS DE SALUD DE GUATEMALA, AREA NOR ORIENTE",
    "DEPARTAMENTO DE BODEGA",
    "CONTROL DE TEMPERATURA MENSUAL AÑO 2026"
]

# Agregar y formatear los títulos del encabezado
for i, titulo in enumerate(titulos_encabezado, start=1):
    ws.merge_cells(f'A{i}:AF{i}')
    cell = ws.cell(row=i, column=1, value=titulo)
    cell.alignment = center_alignment
    cell.font = title_font
    cell.fill = title_fill

# Agregar espacio entre el encabezado y la tabla
current_row = len(titulos_encabezado) + 2

# Títulos de las columnas
headers = ["MES", "HORA"] + [str(i) for i in range(1, 32)]
for col, header in enumerate(headers, 1):
    cell = ws.cell(row=current_row, column=col, value=header)
    cell.fill = header_fill  # Gris oscuro para toda la fila
    cell.font = header_font  # Letra blanca para toda la fila
    cell.alignment = center_alignment
    cell.border = border

months = ["ENERO", "FEBRERO", "MARZO", "ABRIL", "MAYO", "JUNIO",
          "JULIO", "AGOSTO", "SEPTIEMBRE", "OCTUBRE", "NOVIEMBRE", "DICIEMBRE"]

row = current_row + 1
for month_idx, month in enumerate(months, 1):
    # Combinar celdas para el mes
    ws.merge_cells(start_row=row, start_column=1, end_row=row + 1, end_column=1)
    cell = ws.cell(row=row, column=1, value=month)
    cell.alignment = center_alignment
    cell.font = Font(bold=True)
    cell.fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    cell.border = border

    # Agregar AM y PM
    for hora_idx, hora in enumerate(["AM", "PM"]):
        cell = ws.cell(row=row + hora_idx, column=2, value=hora)
        cell.alignment = center_alignment
        cell.fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
        cell.border = border

    # Calcular días para el mes actual
    if month_idx == 12:
        days_in_month = 31
    else:
        days_in_month = (datetime(2026, month_idx + 1, 1) - timedelta(days=1)).day

    # Agregar días y marcar fines de semana y festivos
    for day in range(1, days_in_month + 1):
        current_date = datetime(2026, month_idx, day)
        col = day + 2

        # Aplicar borde a todas las celdas
        ws.cell(row=row, column=col).border = border
        ws.cell(row=row + 1, column=col).border = border

        # Verificar si es festivo
        if (month_idx, day) in festivos_guatemala:
            tipo_festivo, _ = festivos_guatemala[(month_idx, day)]
            valor = tipo_festivo
            ws.cell(row=row, column=col, value=valor).fill = special_fill
            ws.cell(row=row, column=col).font = special_font
            ws.cell(row=row + 1, column=col, value=valor).fill = special_fill
            ws.cell(row=row + 1, column=col).font = special_font
        # Verificar si es fin de semana
        elif current_date.weekday() == 5:  # Sábado
            ws.cell(row=row, column=col, value="S").fill = special_fill
            ws.cell(row=row, column=col).font = special_font
            ws.cell(row=row + 1, column=col, value="S").fill = special_fill
            ws.cell(row=row + 1, column=col).font = special_font
        elif current_date.weekday() == 6:  # Domingo
            ws.cell(row=row, column=col, value="D").fill = special_fill
            ws.cell(row=row, column=col).font = special_font
            ws.cell(row=row + 1, column=col, value="D").fill = special_fill
            ws.cell(row=row + 1, column=col).font = special_font
        else:
            # Dejar las celdas vacías
            ws.cell(row=row, column=col, value="")
            ws.cell(row=row + 1, column=col, value="")

        # Aplicar alineación central a todas las celdas
        ws.cell(row=row, column=col).alignment = center_alignment
        ws.cell(row=row + 1, column=col).alignment = center_alignment

    row += 2

# Ajustar el ancho de las columnas
for col in range(1, ws.max_column + 1):
    max_length = 0
    for row in range(1, ws.max_row + 1):
        cell = ws.cell(row=row, column=col)
        if isinstance(cell, Cell):
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:  # noqa: E722
                pass
    adjusted_width = (max_length + 2)
    ws.column_dimensions[get_column_letter(col)].width = adjusted_width

# Guardar el archivo
wb.save("Calendario_Guatemala_2026.xlsx")
print("✅ Archivo creado: Calendario_Guatemala_2026.xlsx")
print("\n📅 Días festivos incluidos:")
for (mes, dia), (tipo, nombre) in sorted(festivos_guatemala.items()):
    fecha = datetime(2026, mes, dia)
    dia_semana = ["Lunes", "Martes", "Miércoles", "Jueves", "Viernes", "Sábado", "Domingo"][fecha.weekday()]
    print(f"   {dia:2d}/{mes:2d}/2026 ({dia_semana:10s}) - {nombre}")